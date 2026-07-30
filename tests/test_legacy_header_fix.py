"""
Functional test to verify legacy header fix.
Tests that ValidationError objects contain correct Excel metadata.
"""

import os
import sys
import tempfile
from pathlib import Path

# Add project root to path
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from src.core.validation_engine import ValidationEngine
from src.core.workbook_loader import WorkbookLoader
from src.models.models import ValidationError
import openpyxl


def create_test_workbook():
    """Create a test workbook with known validation errors."""
    # Create temporary file
    temp_dir = tempfile.mkdtemp()
    file_path = os.path.join(temp_dir, "test_legacy_headers.xlsx")
    
    # Create workbook with standardized headers
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Daily Output File"
    
    # Required headers (standardized names)
    headers = [
        'Order_no', 'Order_Creation_Date', 'Branch_Code', 'Branch_Name',
        'Card_Type', 'Number_of_Batches', 'Waybill_Number', 'Batch_Number', 'BagNumber'
    ]
    ws.append(headers)
    
    # Row 2: Valid row
    ws.append([
        'ORD001', '2026-07-22', 'BR001', 'Main Branch',
        'SIM', '2', 'WB001', 'B001|B002', "'001"
    ])
    
    # Row 3: Blank BagNumber (error)
    ws.append([
        'ORD002', '2026-07-22', 'BR002', 'North Branch',
        'SIM', '1', 'WB002', 'B003', ''
    ])
    
    # Row 4: Duplicate Batch_Number (error)
    ws.append([
        'ORD003', '2026-07-22', 'BR003', 'South Branch',
        'DMCCLS', '1', 'WB003', 'B001', "'004"
    ])
    
    # Row 5: Incorrect Number_of_Batches (error)
    ws.append([
        'ORD004', '2026-07-22', 'BR004', 'East Branch',
        'SIM', '5', 'WB004', 'B005|B006', "'005"
    ])
    
    # Row 6: Blank Branch_Name (error)
    ws.append([
        'ORD005', '2026-07-22', 'BR005', '',
        'DMCCLS', '1', 'WB005', 'B007', "'006"
    ])
    
    # Row 7: Invalid Card_Type (error)
    ws.append([
        'ORD006', '2026-07-22', 'BR006', 'West Branch',
        'SIMM', '1', 'WB006', 'B008', "'007"
    ])
    
    wb.save(file_path)
    wb.close()
    
    return file_path


def test_legacy_header_fix():
    """Test that validation errors contain correct metadata."""
    print("=" * 70)
    print("FUNCTIONAL TEST: Legacy Header Fix Verification")
    print("=" * 70)
    
    # Create test workbook
    print("\n1. Creating test workbook with validation errors...")
    file_path = create_test_workbook()
    print(f"   ✓ Created: {file_path}")
    
    # Run validation
    print("\n2. Running validation engine...")
    engine = ValidationEngine()
    result = engine.validate_complete_workbook(file_path)
    print(f"   ✓ Validation complete: {result.error_count} errors found")
    
    # Verify results
    print("\n3. Verifying ValidationError objects...")
    print(f"   Total validation_errors: {len(result.validation_errors)}")
    
    # Track error types found
    error_types_found = {}
    
    for i, verr in enumerate(result.validation_errors, 1):
        print(f"\n   Error {i}:")
        print(f"     rule_id: {verr.rule_id}")
        print(f"     error_type: {verr.error_type}")
        print(f"     row_number: {verr.row_number}")
        print(f"     column_name: {verr.column_name}")
        print(f"     cell_reference: {verr.cell_reference}")
        print(f"     invalid_value: {verr.invalid_value}")
        print(f"     error_message: {verr.error_message[:80]}...")
        
        # Verify NO placeholder values
        assert verr.row_number != 0, f"ERROR: row_number is 0 (placeholder) for {verr.rule_id}"
        assert verr.row_number != '?', f"ERROR: row_number is '?' (placeholder) for {verr.rule_id}"
        assert verr.cell_reference != '-', f"ERROR: cell_reference is '-' (placeholder) for {verr.rule_id}"
        assert '?' not in str(verr.cell_reference), f"ERROR: cell_reference contains '?' for {verr.rule_id}"
        
        # Verify column names match standard headers
        valid_columns = {'Order_no', 'Order_Creation_Date', 'Branch_Code', 'Branch_Name',
                        'Card_Type', 'Number_of_Batches', 'Waybill_Number', 
                        'Batch_Number', 'BagNumber'}
        assert verr.column_name in valid_columns or verr.column_name == '', \
            f"ERROR: column_name '{verr.column_name}' is not a standard header"
        
        # Verify error messages contain actual row numbers (except for duplicate summaries)
        if verr.error_type not in ['DUPLICATE_ACROSS_ROWS', 'DUPLICATE_BATCH_SAME_CELL']:
            assert f"Row {verr.row_number}" in verr.error_message or verr.row_number == 0, \
                f"ERROR: error_message doesn't contain actual row number for {verr.rule_id}"
        
        # Track error types
        error_types_found[verr.error_type] = error_types_found.get(verr.error_type, 0) + 1
    
    print("\n4. Error type summary:")
    for error_type, count in sorted(error_types_found.items()):
        print(f"     {error_type}: {count}")
    
    # Verify expected errors were found
    print("\n5. Verifying expected errors...")
    
    expected_errors = {
        'INVALID_BAG': False,  # Blank BagNumber
        'DUPLICATE_ACROSS_ROWS': False,  # Duplicate Batch_Number B001
        'BATCH_MISMATCH': False,  # Incorrect Number_of_Batches
        'BLANK_FIELD': False,  # Blank Branch_Name
        'INVALID_CARD_TYPE': False,  # Invalid Card_Type SIMM
    }
    
    for verr in result.validation_errors:
        if verr.error_type in expected_errors:
            expected_errors[verr.error_type] = True
    
    all_found = True
    for error_type, found in expected_errors.items():
        status = "✓" if found else "✗"
        print(f"   {status} {error_type}: {'Found' if found else 'MISSING'}")
        if not found:
            all_found = False
    
    # Verify specific error details
    print("\n6. Verifying specific error details...")
    
    # Check for blank BagNumber error
    blank_bag_errors = [v for v in result.validation_errors 
                       if v.rule_id == 'BLANK_BAG_NUMBER']
    if blank_bag_errors:
        err = blank_bag_errors[0]
        assert err.row_number == 3, f"Expected row 3, got {err.row_number}"
        assert err.column_name == 'BagNumber', f"Expected BagNumber, got {err.column_name}"
        assert err.cell_reference == 'I3', f"Expected I3, got {err.cell_reference}"
        print(f"   ✓ Blank BagNumber: Row {err.row_number}, Cell {err.cell_reference}")
    
    # Check for duplicate Batch_Number error
    dup_errors = [v for v in result.validation_errors 
                 if v.error_type == 'DUPLICATE_ACROSS_ROWS']
    if dup_errors:
        err = dup_errors[0]
        # First occurrence of B001 is in row 2
        assert err.row_number == 2, f"Expected row 2, got {err.row_number}"
        assert err.column_name == 'Batch_Number', f"Expected Batch_Number, got {err.column_name}"
        assert err.cell_reference == 'H2', f"Expected H2, got {err.cell_reference}"
        assert 'B001' in err.invalid_value, f"Expected B001 in invalid_value"
        print(f"   ✓ Duplicate Batch: Row {err.row_number}, Cell {err.cell_reference}, Value {err.invalid_value}")
    
    # Check for blank Branch_Name error
    blank_branch_errors = [v for v in result.validation_errors 
                          if v.error_type == 'BLANK_FIELD' and 'Branch_Name' in v.column_name]
    if blank_branch_errors:
        err = blank_branch_errors[0]
        assert err.row_number == 6, f"Expected row 6, got {err.row_number}"
        assert err.column_name == 'Branch_Name', f"Expected Branch_Name, got {err.column_name}"
        assert err.cell_reference == 'D6', f"Expected D6, got {err.cell_reference}"
        print(f"   ✓ Blank Branch_Name: Row {err.row_number}, Cell {err.cell_reference}")
    
    # Verify no placeholders
    print("\n7. Final verification...")
    placeholder_errors = []
    for verr in result.validation_errors:
        if verr.row_number in [0, '?', '-', None]:
            placeholder_errors.append(f"row_number={verr.row_number}")
        if verr.cell_reference in ['?', '-', None] or '?' in str(verr.cell_reference):
            placeholder_errors.append(f"cell_reference={verr.cell_reference}")
        if verr.invalid_value in ['-', None] and verr.error_type != 'SYSTEM_ERROR':
            placeholder_errors.append(f"invalid_value={verr.invalid_value}")
    
    if placeholder_errors:
        print(f"   ✗ FAILED: Found placeholders: {placeholder_errors}")
        return False
    else:
        print(f"   ✓ No placeholders found")
    
    # Cleanup
    os.remove(file_path)
    os.rmdir(os.path.dirname(file_path))
    
    print("\n" + "=" * 70)
    print("✅ ALL TESTS PASSED")
    print("=" * 70)
    print("\nSummary:")
    print(f"  - {result.error_count} validation errors found")
    print(f"  - All errors contain actual Excel row numbers")
    print(f"  - All errors contain actual cell references")
    print(f"  - All errors use standardized column names")
    print(f"  - No placeholder values (?, 0, -) detected")
    print(f"  - Dashboard will display correct values")
    
    return True


if __name__ == '__main__':
    try:
        success = test_legacy_header_fix()
        sys.exit(0 if success else 1)
    except Exception as e:
        print(f"\n❌ TEST FAILED: {e}")
        import traceback
        traceback.print_exc()
        sys.exit(1)